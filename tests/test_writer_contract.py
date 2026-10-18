"""Contract tests: the ``parts`` (method A) and ``zip`` (fallback F1) writers.

docs/arch/04 §6 requires the two writers to honour **the same external
contract**: the produced ``.pptx`` opens, keeps every background, lists exactly
``N`` slides, and never duplicates a media part. Any divergence is a bug.
"""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

from lyrics2pptx.errors import ConfigError
from lyrics2pptx.pptx.builder import build_presentation
from lyrics2pptx.pptx.clone import RT_NOTES_SLIDE
from lyrics2pptx.pptx.template_reader import find_title_shape, get_template_profile
from lyrics2pptx.pptx.writer import (
    PartsWriter,
    PptxWriter,
    ZipWriter,
    select_writer,
    write_via_parts,
    write_via_zip,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"
PAIRS = [("하나", "둘"), ("셋", "넷"), ("다섯", None)]


def media_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if name.startswith("ppt/media/")}


def texts_of(path: Path) -> list[list[str]]:
    presentation = Presentation(str(path))
    return [
        [p.text for p in find_title_shape(slide).text_frame.paragraphs]
        for slide in presentation.slides
    ]


class SelectWriterTests(unittest.TestCase):
    def test_forced_strategies_map_to_their_writer(self) -> None:
        self.assertIsInstance(select_writer("parts"), PartsWriter)
        self.assertIsInstance(select_writer("zip"), ZipWriter)

    def test_auto_resolves_to_one_of_the_two_writers(self) -> None:
        self.assertIn(select_writer("auto").name, ("parts", "zip"))

    def test_unknown_strategy_raises_config_error(self) -> None:
        with self.assertRaises(ConfigError):
            select_writer("banana")

    def test_writers_satisfy_the_protocol_and_expose_a_name(self) -> None:
        for writer in (PartsWriter(), ZipWriter()):
            with self.subTest(writer=writer.name):
                self.assertIsInstance(writer, PptxWriter)
                self.assertEqual(writer.name, writer.name.strip())


class WriterEquivalenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        cls.parts_out = root / "parts.pptx"
        cls.zip_out = root / "zip.pptx"
        cls.parts_result = build_presentation(PAIRS, TEMPLATE, cls.parts_out, strategy="parts")
        cls.zip_result = build_presentation(PAIRS, TEMPLATE, cls.zip_out, strategy="zip")

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def test_reported_metadata_is_identical(self) -> None:
        for attribute in ("slide_count", "pool_size", "wrap_count", "used_bg_indices"):
            with self.subTest(attribute=attribute):
                self.assertEqual(
                    getattr(self.parts_result, attribute),
                    getattr(self.zip_result, attribute),
                )

    def test_generated_texts_are_identical(self) -> None:
        self.assertEqual(texts_of(self.parts_out), texts_of(self.zip_out))

    def test_texts_match_the_plan_with_a_one_line_tail(self) -> None:
        self.assertEqual(texts_of(self.parts_out), [["하나", "둘"], ["셋", "넷"], ["다섯"]])

    def test_both_outputs_reopen_with_the_same_slide_count(self) -> None:
        parts = Presentation(str(self.parts_out))
        zipped = Presentation(str(self.zip_out))
        self.assertEqual(len(parts.slides), 3)
        self.assertEqual(len(zipped.slides), len(parts.slides))

    def test_neither_writer_creates_new_media_parts(self) -> None:
        with zipfile.ZipFile(TEMPLATE) as archive:
            template_media = {name for name in archive.namelist() if name.startswith("ppt/media/")}
        self.assertEqual(media_names(self.parts_out), template_media)
        self.assertEqual(media_names(self.zip_out), template_media)

    def test_neither_writer_clones_notes_slides(self) -> None:
        for path in (self.parts_out, self.zip_out):
            with self.subTest(path=path.name):
                presentation = Presentation(str(path))
                for slide in presentation.slides:
                    reltypes = {rel.reltype for rel in slide.part.rels.values()}
                    self.assertNotIn(RT_NOTES_SLIDE, reltypes)

    def test_both_outputs_keep_a_background_reference(self) -> None:
        for path in (self.parts_out, self.zip_out):
            presentation = Presentation(str(path))
            for index, slide in enumerate(presentation.slides):
                with self.subTest(path=path.name, slide=index):
                    xml = etree.tostring(slide._element, encoding="unicode")
                    self.assertTrue(
                        "<p:bg" in xml or "r:embed" in xml,
                        f"{path.name} slide {index} lost its background reference",
                    )


class WriterGuardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls.profile = get_template_profile(TEMPLATE)

    def test_both_writers_reject_an_empty_plan(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "x.pptx"
            with self.assertRaises(ConfigError):
                write_via_parts([], self.profile, out)
            with self.assertRaises(ConfigError):
                write_via_zip([], self.profile, out)


if __name__ == "__main__":
    unittest.main()
