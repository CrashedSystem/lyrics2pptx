"""Integration tests for the slide-clone engine (``pptx/clone.py``).

Uses the real ``template.pptx`` **read-only**. The invariants pinned here come
from docs/arch/04 §3: media parts are reused (never recreated), the copied XML
keeps its background reference, notes are not cloned, and an unmapped
relationship reference is a hard error instead of a silently missing background.
"""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from lyrics2pptx.errors import TemplateStructureError
from lyrics2pptx.pptx.builder import build_presentation
from lyrics2pptx.pptx.clone import (
    RT_NOTES_SLIDE,
    clone_slide,
    remap_relationships,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"


def media_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(1 for name in archive.namelist() if name.startswith("ppt/media/"))


def reltype_set(part) -> set[str]:  # noqa: ANN001 - pptx part proxy
    return {rel.reltype for rel in part.rels.values()}


class CloneSlideTests(unittest.TestCase):
    def setUp(self) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        self.presentation = Presentation(str(TEMPLATE))
        self.slides = list(self.presentation.slides)  # materialize exactly once

    def test_clone_reuses_the_source_media_parts(self) -> None:
        source = self.slides[19]  # full-bleed <p:pic> background
        clone = clone_slide(source, self.presentation)

        source_images = sorted(
            rel.target_part.partname for rel in source.part.rels.values() if rel.reltype == RT.IMAGE
        )
        clone_images = sorted(
            rel.target_part.partname for rel in clone.rels.values() if rel.reltype == RT.IMAGE
        )
        self.assertTrue(clone_images, "clone has no image relationship")
        self.assertEqual(clone_images, source_images, "media parts were recreated instead of reused")

    def test_clone_keeps_the_layout_relationship_and_drops_notes(self) -> None:
        source = self.slides[19]
        clone = clone_slide(source, self.presentation)
        reltypes = reltype_set(clone)
        self.assertIn(RT.SLIDE_LAYOUT, reltypes)
        self.assertNotIn(RT_NOTES_SLIDE, reltypes)
        self.assertIn(RT.NOTES_SLIDE, reltype_set(source.part))

    def test_clone_keeps_notes_when_explicitly_requested(self) -> None:
        clone = clone_slide(self.slides[19], self.presentation, keep_notes=True)
        self.assertIn(RT.NOTES_SLIDE, reltype_set(clone))

    def test_copied_element_keeps_its_background_markup(self) -> None:
        clone = clone_slide(self.slides[0], self.presentation)
        xml = etree.tostring(clone._element, encoding="unicode")
        self.assertIn("<p:bg", xml)

    def test_full_bleed_picture_reference_is_remapped_to_a_valid_rid(self) -> None:
        source = self.slides[19]
        clone = clone_slide(source, self.presentation)
        embeds = [element.get(qn("r:embed")) for element in clone._element.iter() if element.get(qn("r:embed"))]
        self.assertTrue(embeds, "slide20 clone lost its r:embed reference")
        valid = {rel.rId for rel in clone.rels.values()}
        for rid in embeds:
            self.assertIn(rid, valid)

    def test_saved_package_does_not_duplicate_media(self) -> None:
        clone_slide(self.slides[19], self.presentation)
        clone_slide(self.slides[19], self.presentation)
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "clone.pptx"
            self.presentation.save(str(out))
            self.assertEqual(media_count(out), 20)

    def test_unmapped_relationship_reference_is_a_hard_error(self) -> None:
        source = self.slides[19]
        clone = clone_slide(source, self.presentation)
        target = next(
            element for element in clone._element.iter() if element.get(qn("r:embed")) is not None
        )
        target.set(qn("r:embed"), "rId424242")
        with self.assertRaises(TemplateStructureError):
            remap_relationships(source.part, clone)


class CloneInsideBuildTests(unittest.TestCase):
    """The clone engine as the builder actually drives it."""

    def test_single_slide_build_from_a_full_bleed_template_slide(self) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "one.pptx"
            result = build_presentation(
                [("첫째 줄", "둘째 줄")],
                TEMPLATE,
                out,
                slide_index=20,  # template slide 20 uses a full-bleed <p:pic>
                strategy="parts",
            )
            self.assertEqual(result.slide_index, 20)
            # Only the selected background's media survives orphan pruning.
            self.assertEqual(media_count(out), 1)

            presentation = Presentation(str(out))
            self.assertEqual(len(presentation.slides), 1)
            slide = list(presentation.slides)[0]
            reltypes = reltype_set(slide.part)
            self.assertIn(RT.IMAGE, reltypes, "background picture relationship missing")
            self.assertNotIn(RT.NOTES_SLIDE, reltypes, "notes slide was cloned")
            xml = etree.tostring(slide._element, encoding="unicode")
            self.assertTrue("<p:bg" in xml or "r:embed" in xml)


if __name__ == "__main__":
    unittest.main()
