"""Integration + boundary tests for ``pptx/builder.py``.

Covers the two rules from docs/arch/04 §2:

* ``bg_index = pool[k % P]`` including the wrap-around cases
  (0/1/2/20/21/40/41/42 lines);
* the "one usable line" verdict (``EmptyLyricsError``) and the reopen check.

Real builds use the read-only ``template.pptx``; every produced file is written
to a temporary directory and reopened with python-pptx.
"""

from __future__ import annotations

import dataclasses
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from pptx import Presentation
from pptx.oxml.ns import qn

from lyrics2pptx import config
from lyrics2pptx.errors import ConfigError, EmptyLyricsError
from lyrics2pptx.pptx.builder import build_plan, build_presentation
from lyrics2pptx.pptx.clone import RT_NOTES_SLIDE
from lyrics2pptx.pptx.template_reader import (
    get_template_profile,
    find_title_shape,
)
from lyrics2pptx.text.split import split_pairs

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"


def media_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(1 for name in archive.namelist() if name.startswith("ppt/media/"))


def slide_member_targets(path: Path) -> list[str]:
    """Return the slide part targets referenced by ``<p:sldIdLst>``, in order."""
    from lxml import etree

    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(path) as archive:
        presentation = etree.fromstring(archive.read("ppt/presentation.xml"))
        rels = etree.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
    id_to_target = {
        rel.get("Id"): rel.get("Target")
        for rel in rels.findall(f"{{{rel_ns}}}Relationship")
    }
    targets: list[str] = []
    for sld_id in presentation.iter(f"{{{p_ns}}}sldId"):
        targets.append(id_to_target[sld_id.get(f"{{{r_ns}}}id")])
    return targets


def slide_xml(path: Path, index: int) -> str:
    """Return the raw XML of the ``index``-th slide in ``sldIdLst`` order."""
    target = slide_member_targets(path)[index]
    with zipfile.ZipFile(path) as archive:
        return archive.read("ppt/" + target.lstrip("/")).decode("utf-8")


def paragraph_texts(path: Path) -> list[list[str]]:
    presentation = Presentation(str(path))
    texts: list[list[str]] = []
    for slide in presentation.slides:
        shape = find_title_shape(slide)
        texts.append([p.text for p in shape.text_frame.paragraphs])
    return texts


def slide_reltypes(path: Path) -> list[set[str]]:
    presentation = Presentation(str(path))
    return [{rel.reltype for rel in slide.part.rels.values()} for slide in presentation.slides]


class PlanBoundaryTests(unittest.TestCase):
    """The cycling formula, exercised without touching the disk."""

    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls.profile = get_template_profile(TEMPLATE)

    @staticmethod
    def _pairs(line_count: int) -> list[tuple[str, str | None]]:
        return split_pairs([f"line{k:02d}" for k in range(line_count)])

    def test_slide_count_equals_ceil_of_half_the_line_count(self) -> None:
        expected = {0: 0, 1: 1, 2: 1, 3: 2, 20: 10, 21: 11, 40: 20, 41: 21, 42: 21}
        for line_count, slides in expected.items():
            with self.subTest(lines=line_count):
                self.assertEqual(len(self._pairs(line_count)), slides)

    def test_background_index_is_pool_index_modulo_pool_size(self) -> None:
        for line_count in (1, 2, 3, 20, 21, 40, 41, 42):
            with self.subTest(lines=line_count):
                pairs = self._pairs(line_count)
                specs = build_plan(pairs, self.profile)
                pool = self.profile.pool
                self.assertEqual(
                    [spec.bg_index for spec in specs],
                    [pool[k % len(pool)] for k in range(len(pairs))],
                )

    def test_wrap_around_examples_from_the_architecture(self) -> None:
        # 42 lines -> 21 slides: the 21st (k == 20) wraps back to pool[0].
        specs42 = build_plan(self._pairs(42), self.profile)
        self.assertEqual(len(specs42), 21)
        self.assertEqual(specs42[20].bg_index, self.profile.pool[0])
        self.assertEqual([spec.bg_index for spec in specs42], [*range(20), 0])

        # 40 lines -> 20 slides: exactly one full cycle, no wrap.
        specs40 = build_plan(self._pairs(40), self.profile)
        self.assertEqual([spec.bg_index for spec in specs40], list(range(20)))

        # 41 lines -> 21 slides: odd tail wraps too.
        specs41 = build_plan(self._pairs(41), self.profile)
        self.assertEqual(len(specs41), 21)
        self.assertEqual(specs41[20].bg_index, self.profile.pool[0])

    def test_style_comes_from_the_same_template_slide_as_the_background(self) -> None:
        specs = build_plan(self._pairs(42), self.profile)
        for spec in specs:
            self.assertIs(spec.style, self.profile.styles[spec.bg_index])

    def test_empty_plan_is_returned_for_no_pairs(self) -> None:
        self.assertEqual(build_plan([], self.profile), [])

    def test_empty_pool_raises_config_error(self) -> None:
        empty = dataclasses.replace(self.profile, pool=())
        with self.assertRaises(ConfigError):
            build_plan(self._pairs(2), empty)

    def test_skip_slide_changes_the_pool_and_therefore_the_cycle(self) -> None:
        profile = get_template_profile(TEMPLATE, skip_slide=(1,))
        specs = build_plan(self._pairs(40), profile)
        self.assertEqual(specs[0].bg_index, 1)
        self.assertEqual([spec.bg_index for spec in specs], list(range(1, 20)) + [1])


class BuilderErrorTests(unittest.TestCase):
    def test_zero_lines_raise_empty_lyrics_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "empty.pptx"
            with self.assertRaises(EmptyLyricsError):
                build_presentation([], TEMPLATE, out)
            self.assertFalse(out.exists())

    def test_unknown_strategy_raises_config_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "bad.pptx"
            with self.assertRaises(ConfigError):
                build_presentation([("a", "b")], TEMPLATE, out, strategy="banana")
            self.assertFalse(out.exists())

    def test_missing_template_raises_config_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "x.pptx"
            with self.assertRaises(ConfigError):
                build_presentation([("a", "b")], Path(tmp) / "nope.pptx", out)

    def test_atomic_replace_leaves_no_temp_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "ok.pptx"
            build_presentation([("a", "b")], TEMPLATE, out)
            leftovers = [p.name for p in Path(tmp).iterdir() if p.name.startswith(".tmp-")]
            self.assertEqual(leftovers, [])


class BuilderIntegrationTests(unittest.TestCase):
    """Real builds for the boundary line counts, each reopened afterwards."""

    CASES = {"42": 42, "40": 40, "3": 3}

    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls._tmp = tempfile.TemporaryDirectory()
        root = Path(cls._tmp.name)
        cls.built: dict[str, dict] = {}
        for key, line_count in cls.CASES.items():
            lines = [f"line{k:02d}" for k in range(line_count)]
            pairs = split_pairs(lines)
            out = root / f"{key}.pptx"
            result = build_presentation(pairs, TEMPLATE, out, strategy="auto")
            cls.built[key] = {"lines": lines, "pairs": pairs, "out": out, "result": result}

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def test_42_lines_produce_21_slides_with_one_wrap(self) -> None:
        case = self.built["42"]
        result = case["result"]
        self.assertEqual(result.slide_count, 21)
        self.assertEqual(result.pool_size, 20)
        self.assertEqual(result.wrap_count, 1)
        self.assertEqual(result.used_bg_indices, tuple([*range(20), 0]))
        self.assertEqual(result.dropped_notes_count, 21)

    def test_40_lines_produce_20_slides_with_no_wrap(self) -> None:
        result = self.built["40"]["result"]
        self.assertEqual(result.slide_count, 20)
        self.assertEqual(result.pool_size, 20)
        self.assertEqual(result.wrap_count, 0)
        self.assertEqual(result.used_bg_indices, tuple(range(20)))
        self.assertEqual(result.dropped_notes_count, 20)

    def test_3_lines_produce_2_slides_with_a_one_line_tail(self) -> None:
        result = self.built["3"]["result"]
        self.assertEqual(result.slide_count, 2)
        self.assertEqual(result.used_bg_indices, (0, 1))

    def test_output_reopens_and_lists_exactly_the_built_slides(self) -> None:
        for key, case in self.built.items():
            with self.subTest(case=key):
                presentation = Presentation(str(case["out"]))
                self.assertEqual(len(presentation.slides), len(case["pairs"]))

    def test_each_slide_carries_its_two_lines(self) -> None:
        case = self.built["42"]
        texts = paragraph_texts(case["out"])
        self.assertEqual(len(texts), 21)
        for index, (first, second) in enumerate(case["pairs"]):
            self.assertEqual(texts[index], [first, second])

    def test_the_odd_tail_slide_carries_exactly_one_line(self) -> None:
        texts = paragraph_texts(self.built["3"]["out"])
        self.assertEqual(texts[-1], ["line02"])

    def test_media_parts_are_not_duplicated(self) -> None:
        for key, case in self.built.items():
            with self.subTest(case=key):
                self.assertEqual(media_count(case["out"]), 20)

    def test_every_slide_keeps_a_background_reference(self) -> None:
        case = self.built["42"]
        for index in range(case["result"].slide_count):
            with self.subTest(slide=index):
                xml = slide_xml(case["out"], index)
                self.assertTrue(
                    "<p:bg" in xml or "r:embed" in xml,
                    f"slide {index} lost its background reference",
                )

    def test_notes_slides_are_not_cloned(self) -> None:
        for key, case in self.built.items():
            with self.subTest(case=key):
                for reltypes in slide_reltypes(case["out"]):
                    self.assertNotIn(RT_NOTES_SLIDE, reltypes)

    def test_sld_id_list_length_equals_the_slide_count(self) -> None:
        case = self.built["42"]
        presentation = Presentation(str(case["out"]))
        sld_id_lst = presentation.part._element.find(qn("p:sldIdLst"))
        self.assertIsNotNone(sld_id_lst)
        self.assertEqual(len(sld_id_lst.findall(qn("p:sldId"))), 21)

    def test_default_output_path_honours_config_output_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            with mock.patch.object(config, "OUTPUT_DIR", Path(tmp)):
                result = build_presentation([("a", "b")], TEMPLATE, None)
            self.assertEqual(result.output_path.parent, Path(tmp))
            self.assertTrue(result.output_path.is_file())
            self.assertEqual(result.output_path.name, config.DEFAULT_FILENAME)


if __name__ == "__main__":
    unittest.main()
