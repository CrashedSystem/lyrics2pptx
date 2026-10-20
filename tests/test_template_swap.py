"""Template swap: a replaced ``template.pptx`` must just work (ADR-034).

The bundled template names its lyric shape ``제목 1`` and gives it a styled run.
A template made elsewhere usually does neither: the shape is called something
else, and a "background-only" template has an empty text box with no run
formatting at all. Before ADR-034 both cases ended in ``TemplateStructureError``.

These tests build small synthetic templates with python-pptx (no 37 MB fixture)
and pin the three promises:

* the shape is found by name, else guessed deterministically;
* a missing run style is synthesized from the box height (``normAutofit`` still
  protects the render);
* a solid-colour background still produces a thumbnail, so the web UI grid is
  never blank after a swap.
"""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from lyrics2pptx import config, pipeline
from lyrics2pptx.errors import ConfigError, TemplateStructureError
from lyrics2pptx.pptx.builder import build_presentation
from lyrics2pptx.pptx.template_reader import clear_profile_cache, get_template_profile
from lyrics2pptx.pptx.thumbnails import (
    clear_thumbnail_cache,
    get_slide_thumbnail,
    list_slide_meta,
)

SLIDE_W = 9144000  # 10 in — 4:3, deliberately not the bundled 16:9
SLIDE_H = 6858000  # 7.5 in
BOX = (Inches(0.9), Inches(2.0), Inches(7.6), Inches(1.5))
SOLID_RGB = (0x11, 0x22, 0x33)


def write_template(
    path: Path,
    *,
    shape_name: str = "제목 1",
    with_runs: bool = True,
    slides: int = 3,
    extra_box: bool = False,
    picture_background: Path | None = None,
    background_shape: bool = True,
) -> Path:
    """Build a tiny 4:3 template and return its path."""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    for index in range(slides):
        slide = prs.slides.add_slide(blank)
        if picture_background is not None:
            slide.shapes.add_picture(
                str(picture_background), 0, 0, width=prs.slide_width,
                height=prs.slide_height,
            )
        elif background_shape:
            rect = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
            )
            rect.name = "배경 사각형"
            rect.fill.solid()
            rect.fill.fore_color.rgb = RGBColor(*SOLID_RGB)
            rect.line.fill.background()

        box = slide.shapes.add_textbox(*BOX)
        box.name = shape_name
        if with_runs:
            paragraph = box.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = "가사 자리"
            run.font.size = Pt(40)
            run.font.name = "맑은 고딕"
            paragraph.alignment = PP_ALIGN.CENTER

        if extra_box:
            big = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9.0), Inches(2.6))
            big.name = "큰 텍스트 상자"
            big.text_frame.paragraphs[0].add_run().text = "가사가 들어갈 자리"

    prs.save(str(path))
    return path


class _TemplateDirTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        clear_profile_cache()
        clear_thumbnail_cache()
        self.addCleanup(clear_profile_cache)
        self.addCleanup(clear_thumbnail_cache)


class DetectedShapeBuildTests(_TemplateDirTestCase):
    def test_another_shape_name_builds_with_the_synthesized_size(self) -> None:
        template = write_template(
            self.tmp / "other-name.pptx", shape_name="Title A", with_runs=False
        )
        out = self.tmp / "out.pptx"

        result = build_presentation(
            [("첫째 줄", "둘째 줄"), ("셋째 줄", "넷째 줄")],
            template,
            out,
            slide_index=1,
        )

        self.assertEqual(result.slide_count, 2)
        produced = Presentation(str(out))
        self.assertEqual((produced.slide_width, produced.slide_height), (SLIDE_W, SLIDE_H))

        shape = next(s for s in produced.slides[0].shapes if s.name == "Title A")
        self.assertEqual(
            (shape.left, shape.top, shape.width, shape.height), tuple(BOX)
        )
        paragraphs = shape.text_frame.paragraphs
        self.assertEqual([p.text for p in paragraphs], ["첫째 줄", "둘째 줄"])
        # 1.5 in = 108 pt -> 108 * 0.42 = 45 pt (ADR-034 estimate).
        self.assertEqual(shape.text_frame.paragraphs[0].runs[0].font.size, Pt(45))
        body_pr = shape._element.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
        self.assertIsNotNone(body_pr.find(qn("a:normAutofit")))

    def test_the_largest_text_shape_is_chosen(self) -> None:
        template = write_template(
            self.tmp / "two-boxes.pptx", shape_name="작은 상자", extra_box=True
        )
        out = self.tmp / "out.pptx"

        build_presentation([("가", "나")], template, out, slide_index=1)

        produced = Presentation(str(out))
        big = next(s for s in produced.slides[0].shapes if s.name == "큰 텍스트 상자")
        small = next(s for s in produced.slides[0].shapes if s.name == "작은 상자")
        self.assertEqual([p.text for p in big.text_frame.paragraphs], ["가", "나"])
        self.assertEqual([p.text for p in small.text_frame.paragraphs], ["가사 자리"])

    def test_a_shape_named_title_1_keeps_the_template_style(self) -> None:
        template = write_template(self.tmp / "named.pptx")
        profile = get_template_profile(template, slide_index=1)

        self.assertEqual(profile.selected_shape_name, "제목 1")
        self.assertFalse(profile.shape_detected)
        self.assertFalse(profile.styles[0].synthesized)
        self.assertEqual(profile.styles[0].sz, 4000)

    def test_a_background_only_slide_raises_with_the_available_names(self) -> None:
        prs = Presentation()
        prs.slide_width = Emu(SLIDE_W)
        prs.slide_height = Emu(SLIDE_H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        rect = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        rect.name = "배경만"
        path = self.tmp / "no-text.pptx"
        prs.save(str(path))

        with self.assertRaises(TemplateStructureError) as ctx:
            get_template_profile(path, slide_index=1)
        message = str(ctx.exception)
        self.assertIn("배경만", message)
        self.assertIn("LYRICS2PPTX_TITLE_SHAPE", message)

    def test_the_default_slide_number_falls_back_to_the_last_slide(self) -> None:
        template = write_template(self.tmp / "three.pptx", slides=3)

        profile = get_template_profile(template)  # default 19 > 3

        self.assertEqual(profile.selected_slide, 2)
        self.assertEqual(profile.slide_count, 3)

    def test_resolve_slide_index_reports_the_fallback(self) -> None:
        template = write_template(self.tmp / "three.pptx", slides=3)

        with mock.patch.object(config, "TEMPLATE_PATH", template):
            self.assertEqual(pipeline.resolve_slide_index(None), 3)
            self.assertEqual(pipeline.resolve_slide_index(2), 2)
            with self.assertRaises(ConfigError):
                pipeline.resolve_slide_index(9)

    def test_an_explicit_out_of_range_slide_is_still_rejected(self) -> None:
        template = write_template(self.tmp / "three.pptx", slides=3)

        with self.assertRaises(ConfigError):
            get_template_profile(template, slide_index=5)


class SolidBackgroundThumbnailTests(_TemplateDirTestCase):
    def test_a_solid_background_is_painted_as_png(self) -> None:
        template = write_template(self.tmp / "solid.pptx", shape_name="Title A", with_runs=False)

        thumbnail = get_slide_thumbnail(1, template)

        self.assertEqual(thumbnail.content_type, "image/png")
        image = Image.open(io.BytesIO(thumbnail.data))
        self.assertEqual(image.size, (640, 480))  # 4:3, 640 px wide
        self.assertEqual(image.convert("RGB").getpixel((10, 10)), SOLID_RGB)

    def test_a_picture_background_is_served_unchanged(self) -> None:
        source = self.tmp / "bg.png"
        Image.new("RGB", (64, 48), (200, 100, 50)).save(source)
        template = write_template(
            self.tmp / "picture.pptx",
            shape_name="Title A",
            with_runs=False,
            picture_background=source,
            background_shape=False,
        )

        thumbnail = get_slide_thumbnail(1, template)

        self.assertEqual(thumbnail.data, source.read_bytes())

    def test_meta_reports_the_detected_shape(self) -> None:
        template = write_template(self.tmp / "solid.pptx", shape_name="Title A", with_runs=False)

        meta = list_slide_meta(template)

        self.assertEqual(len(meta), 3)
        self.assertTrue(all(item.has_title_shape for item in meta))


class DoctorReportTests(_TemplateDirTestCase):
    def test_doctor_names_the_detected_shape_and_the_style_source(self) -> None:
        template = write_template(
            self.tmp / "other-name.pptx", shape_name="Title A", with_runs=False
        )

        with mock.patch.object(config, "TEMPLATE_PATH", template), mock.patch.object(
            config, "SLIDE_INDEX_IS_EXPLICIT", False
        ):
            report = pipeline.doctor_report()

        self.assertEqual(report["title_shape"], "Title A")
        self.assertTrue(report["title_shape_detected"])
        self.assertEqual(report["style_source"], "synthesized")
        self.assertEqual(report["selected_slide"], 3)  # default falls back to the last slide


if __name__ == "__main__":
    unittest.main()
