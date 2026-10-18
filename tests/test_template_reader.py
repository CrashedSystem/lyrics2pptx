"""Integration tests for ``pptx/template_reader`` against the real template.

The template is opened read-only (docs/arch/05 §6). These tests pin the measured
properties the whole generator depends on: 20 slides, a 20-entry background pool,
16:9 slide size, and a ``제목 1`` style snapshot per slide.
"""

from __future__ import annotations

import unittest
from pathlib import Path

from pptx import Presentation

from lyrics2pptx.errors import ConfigError, TemplateStructureError
from lyrics2pptx.pptx.template_reader import (
    clear_profile_cache,
    extract_slide_style,
    find_title_shape,
    find_title_shape_element,
    get_template_profile,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"


class TemplateProfileTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls.profile = get_template_profile(TEMPLATE)

    def test_template_has_twenty_slides(self) -> None:
        self.assertEqual(self.profile.slide_count, 20)

    def test_default_background_pool_is_every_slide_in_order(self) -> None:
        self.assertEqual(self.profile.pool, tuple(range(20)))

    def test_slide_size_is_16_by_9(self) -> None:
        self.assertEqual(self.profile.slide_size, (12192000, 6858000))

    def test_template_size_matches_the_file_on_disk(self) -> None:
        self.assertEqual(self.profile.template_size, TEMPLATE.stat().st_size)

    def test_every_slide_exposes_the_title_shape(self) -> None:
        missing = [
            index + 1
            for index, style in sorted(self.profile.styles.items())
            if not style.shape_found
        ]
        self.assertEqual(missing, [])

    def test_every_slide_has_a_style_snapshot_but_the_sizes_vary(self) -> None:
        sizes = {style.sz for style in self.profile.styles.values()}
        self.assertTrue(all(size is not None for size in sizes))
        self.assertTrue(all(6000 <= size <= 7200 for size in sizes), sizes)
        # 계획서 §글꼴 크기: 슬라이드마다 다름 (slide19 = 6600, 나머지 7200).
        self.assertEqual(self.profile.styles[18].sz, 6600)
        self.assertEqual(self.profile.styles[1].sz, 7200)

    def test_default_style_snapshot_is_available_for_fallback(self) -> None:
        default = self.profile.default_style
        self.assertTrue(default.has_rpr)
        self.assertTrue(default.rpr_xml)
        self.assertEqual(default.slide_index, 0)

    def test_profile_is_cached_for_the_same_path_size_and_mtime(self) -> None:
        # Two consecutive calls must not re-analyze the 37 MB template.
        first = get_template_profile(TEMPLATE)
        second = get_template_profile(TEMPLATE)
        self.assertIs(first, second)

    def test_cache_is_invalidated_by_a_clear(self) -> None:
        clear_profile_cache()
        self.addCleanup(clear_profile_cache)
        rebuilt = get_template_profile(TEMPLATE)
        self.assertEqual(rebuilt.slide_count, 20)


class TemplateValidationTests(unittest.TestCase):
    def test_missing_template_raises_config_error(self) -> None:
        with self.assertRaises(ConfigError):
            get_template_profile(Path("D:/definitely/not/here.pptx"))

    def test_skip_slide_out_of_range_raises_config_error(self) -> None:
        with self.assertRaises(ConfigError):
            get_template_profile(TEMPLATE, skip_slide=(21,))

    def test_skip_slide_below_one_raises_config_error(self) -> None:
        with self.assertRaises(ConfigError):
            get_template_profile(TEMPLATE, skip_slide=(0,))

    def test_skip_slide_removes_the_slide_from_the_pool(self) -> None:
        profile = get_template_profile(TEMPLATE, skip_slide=(1, 20))
        self.assertEqual(profile.pool, tuple(range(1, 19)))
        self.assertEqual(profile.slide_count, 20)


class TitleShapeLookupTests(unittest.TestCase):
    def test_shape_is_found_by_name_on_a_template_slide(self) -> None:
        presentation = Presentation(str(TEMPLATE))
        shape = find_title_shape(presentation.slides[18])
        self.assertEqual(shape.name.strip(), "제목 1")
        self.assertEqual([p.text for p in shape.text_frame.paragraphs], ["우리 함께 손잡고 가세", "조금 느리고 어려울지라도"])

    def test_missing_shape_raises_with_the_available_names(self) -> None:
        presentation = Presentation()  # python-pptx default layout: "Title 1"
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
        with self.assertRaises(TemplateStructureError) as ctx:
            find_title_shape(slide)
        self.assertIn("제목 1", str(ctx.exception))

    def test_element_lookup_returns_none_for_a_blank_slide(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertIsNone(find_title_shape_element(slide._element))


class ExtractStyleTests(unittest.TestCase):
    def test_snapshot_of_slide_19_carries_sz_and_serialized_rpr(self) -> None:
        presentation = Presentation(str(TEMPLATE))
        style = extract_slide_style(presentation.slides[18], 18)
        self.assertEqual(style.slide_index, 18)
        self.assertTrue(style.shape_found)
        self.assertTrue(style.has_rpr)
        self.assertEqual(style.sz, 6600)
        self.assertIn(b"rPr", style.rpr_xml)
        self.assertEqual(style.paragraph_count, 2)


if __name__ == "__main__":
    unittest.main()
