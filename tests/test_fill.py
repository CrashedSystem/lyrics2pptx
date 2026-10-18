"""Integration tests for ``pptx/fill.py`` — the ``제목 1`` reconstruction (F1..F8).

Rules pinned here (docs/arch/04 §4):

* F1 the run formatting (``rPr``/``sz``) is inherited from the background slide,
  so per-slide font sizes survive;
* F2 exactly one paragraph on an odd last slide, two otherwise — never three;
* F4 ``normAutofit`` is always present after filling;
* F6 the shape position and size are untouched;
* F7 text is set through ``a:t`` (escaping handled by lxml).
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from lyrics2pptx.errors import TemplateStructureError
from lyrics2pptx.pptx.builder import build_presentation
from lyrics2pptx.pptx.fill import (
    ensure_norm_autofit,
    fill_title1,
    fill_txBody,
    find_title1,
    has_norm_autofit,
)
from lyrics2pptx.pptx.template_reader import find_title_shape, get_template_profile

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"

# Background pool reduced to template slide 19 (0-based index 18, sz == 6600).
SKIP_TO_SLIDE_19 = tuple([*range(1, 19), 20])

PAIRS = [("A & B < C", "2 > 1"), ("odd line", None)]


def _first_run_sz(paragraph) -> str | None:  # noqa: ANN001 - pptx proxy
    for run in paragraph.runs:
        rpr = run._r.find(qn("a:rPr"))
        if rpr is not None:
            return rpr.get("sz")
    return None


def _shape_box(shape) -> tuple:  # noqa: ANN001 - pptx proxy
    return (shape.left, shape.top, shape.width, shape.height)


class FillInBuildTests(unittest.TestCase):
    """One real build reused for every fill assertion."""

    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls._tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls._tmp.name) / "fill.pptx"
        cls.result = build_presentation(
            PAIRS, TEMPLATE, cls.out, skip_slide=SKIP_TO_SLIDE_19, strategy="parts"
        )
        cls.presentation = Presentation(str(cls.out))
        cls.slides = list(cls.presentation.slides)
        cls.template = Presentation(str(TEMPLATE))
        cls.template_shape = find_title_shape(cls.template.slides[18])

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def test_two_line_slide_has_exactly_two_paragraphs(self) -> None:
        shape = find_title1(self.slides[0])
        self.assertEqual([p.text for p in shape.text_frame.paragraphs], ["A & B < C", "2 > 1"])

    def test_odd_last_slide_has_exactly_one_paragraph(self) -> None:
        shape = find_title1(self.slides[1])
        self.assertEqual([p.text for p in shape.text_frame.paragraphs], ["odd line"])
        self.assertEqual(len(shape.text_frame.paragraphs), 1)

    def test_font_size_is_inherited_from_the_background_slide(self) -> None:
        for slide in self.slides:
            shape = find_title1(slide)
            for paragraph in shape.text_frame.paragraphs:
                self.assertEqual(_first_run_sz(paragraph), "6600")

    def test_norm_autofit_is_present_on_every_filled_slide(self) -> None:
        for slide in self.slides:
            shape = find_title1(slide)
            self.assertTrue(has_norm_autofit(shape.text_frame._txBody))

    def test_shape_position_and_size_are_unchanged(self) -> None:
        for slide in self.slides:
            shape = find_title1(slide)
            self.assertEqual(_shape_box(shape), _shape_box(self.template_shape))

    def test_special_characters_are_escaped_and_round_trip(self) -> None:
        # A literal "&" and "<" must survive as text, not break the XML.
        shape = find_title1(self.slides[0])
        self.assertEqual(shape.text_frame.paragraphs[0].text, "A & B < C")
        self.assertEqual(shape.text_frame.paragraphs[1].text, "2 > 1")

    def test_paragraphs_are_not_fragmented_into_extra_runs(self) -> None:
        for slide in self.slides:
            shape = find_title1(slide)
            for paragraph in shape.text_frame.paragraphs:
                self.assertEqual(len(paragraph.runs), 1)


class FillUnitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.is_file():
            raise unittest.SkipTest(f"template not found: {TEMPLATE}")
        cls.style = get_template_profile(TEMPLATE).default_style

    def test_fill_tx_body_rejects_an_empty_first_line(self) -> None:
        tx_body = parse_xml(f'<a:txBody {nsdecls("a")}><a:bodyPr/><a:p/></a:txBody>')
        with self.assertRaises(TemplateStructureError):
            fill_txBody(tx_body, "", "second", style=self.style)

    def test_find_title1_raises_when_the_shape_is_absent(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        with self.assertRaises(TemplateStructureError):
            find_title1(slide)

    def test_fill_title1_raises_when_the_shape_is_absent(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        with self.assertRaises(TemplateStructureError):
            fill_title1(slide, "a", "b", style=self.style)


class NormAutofitTests(unittest.TestCase):
    def test_has_norm_autofit_is_false_without_the_element(self) -> None:
        tx_body = parse_xml(f'<a:txBody {nsdecls("a")}><a:bodyPr/><a:p/></a:txBody>')
        self.assertFalse(has_norm_autofit(tx_body))

    def test_ensure_norm_autofit_adds_the_element_and_keeps_attributes(self) -> None:
        body_pr = parse_xml(f'<a:bodyPr {nsdecls("a")} anchor="ctr" wrap="square"/>')
        ensure_norm_autofit(body_pr)
        self.assertIsNotNone(body_pr.find(qn("a:normAutofit")))
        self.assertEqual(body_pr.get("anchor"), "ctr")
        self.assertEqual(body_pr.get("wrap"), "square")

    def test_ensure_norm_autofit_removes_conflicting_autofit_modes(self) -> None:
        body_pr = parse_xml(
            f'<a:bodyPr {nsdecls("a")}><a:spAutoFit/><a:noAutofit/></a:bodyPr>'
        )
        ensure_norm_autofit(body_pr)
        self.assertIsNone(body_pr.find(qn("a:spAutoFit")))
        self.assertIsNone(body_pr.find(qn("a:noAutofit")))
        self.assertIsNotNone(body_pr.find(qn("a:normAutofit")))

    def test_ensure_norm_autofit_is_idempotent(self) -> None:
        body_pr = parse_xml(f'<a:bodyPr {nsdecls("a")}/>')
        ensure_norm_autofit(body_pr)
        ensure_norm_autofit(body_pr)
        self.assertEqual(len(body_pr.findall(qn("a:normAutofit"))), 1)


if __name__ == "__main__":
    unittest.main()
