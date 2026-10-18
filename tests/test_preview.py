"""HTTP integration tests for the preview route ``POST /api/preview`` (M1).

The browser preview must show exactly what a build would put on the slides, so
normalization and two-line pairing live on the server only
(``pipeline.preview`` → ``text.normalize``). These tests drive the real
``ThreadingHTTPServer`` over a socket and compare the preview pairs against the
paragraph text of an actual build of the same input.
"""

from __future__ import annotations

import http.client
import json
import logging
import tempfile
import threading
import unittest
from pathlib import Path
from unittest import mock

from pptx import Presentation

from lyrics2pptx import config, pipeline
from lyrics2pptx.pptx.template_reader import find_title_shape
from lyrics2pptx.webapp import build_server

TEMPLATE = Path(__file__).resolve().parents[1] / "template.pptx"

# Mixture of markers, blanks and real lyric lines that merely start with a
# keyword word ("chorus line", "intro to me") — those must survive (M2).
RAW = (
    "[Verse 1]\n"
    "너는 담장 너머로\n"
    "chorus line\n"
    "\n"
    "[Chorus]\n"
    "붉은 나무가지에 푸른 열매처럼\n"
    "×2\n"
    "intro to me"
)


def setUpModule() -> None:  # noqa: N802 - unittest hook
    logging.disable(logging.CRITICAL)


def tearDownModule() -> None:  # noqa: N802 - unittest hook
    logging.disable(logging.NOTSET)


def slide_paragraph_texts(path: Path) -> list[list[str]]:
    """Return the ``제목 1`` paragraph text of every slide, in slide order."""
    presentation = Presentation(str(path))
    return [
        [p.text for p in find_title_shape(slide).text_frame.paragraphs]
        for slide in presentation.slides
    ]


class PreviewRouteTestCase(unittest.TestCase):
    """Boots the real server on 127.0.0.1 with an ephemeral port."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.out_dir = Path(self._tmp.name)
        patcher = mock.patch.object(config, "OUTPUT_DIR", self.out_dir)
        patcher.start()
        self.addCleanup(patcher.stop)
        pipeline.clear_search_cache()

        self.server = build_server("127.0.0.1", 0)
        self.host = self.server.server_address[0]
        self.port = self.server.server_address[1]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self._shutdown)

    def _shutdown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=10)
        self.assertFalse(self.thread.is_alive(), "server thread did not stop")
        self._tmp.cleanup()
        pipeline.clear_search_cache()

    # -- helpers ----------------------------------------------------------

    def request(self, method: str, path: str, *, body=None, headers=None):
        connection = http.client.HTTPConnection(self.host, self.port, timeout=60)
        try:
            merged = {"Connection": "close"}
            merged.update(headers or {})
            connection.request(method, path, body=body, headers=merged)
            response = connection.getresponse()
            payload = response.read()
            return response.status, {k.lower(): v for k, v in response.getheaders()}, payload
        finally:
            connection.close()

    def post_json(self, path: str, payload) -> tuple:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        return self.request(
            "POST", path, body=body, headers={"Content-Type": "application/json"}
        )

    def preview(self, payload) -> dict:
        status, _, body = self.post_json("/api/preview", payload)
        self.assertEqual(status, 200, body)
        return json.loads(body)


class PreviewMatchesBuildTests(PreviewRouteTestCase):
    """M1: the preview path and the build path must agree, line for line."""

    def setUp(self) -> None:
        super().setUp()
        if not TEMPLATE.is_file():
            self.skipTest(f"template not found: {TEMPLATE}")

    def test_preview_pairs_equal_the_generated_slide_text(self) -> None:
        preview = self.preview({"raw_text": RAW})
        expected = [[a] if b is None else [a, b] for a, b in preview["pairs"]]

        status, _, pptx = self.post_json(
            "/api/generate", {"raw_text": RAW, "filename": "preview-check.pptx"}
        )
        self.assertEqual(status, 200)
        self.assertTrue(pptx.startswith(b"PK"), "response is not a zip/pptx payload")

        built = slide_paragraph_texts(self.out_dir / "preview-check.pptx")
        self.assertTrue(built, "the build produced no slides")
        self.assertEqual(built, expected)

    def test_preview_lines_are_the_flat_pairs(self) -> None:
        preview = self.preview({"raw_text": RAW})
        flat = [line for pair in preview["pairs"] for line in pair if line is not None]
        self.assertEqual(preview["lines"], flat)

    def test_preview_does_not_write_a_pptx(self) -> None:
        self.preview({"raw_text": RAW})
        self.assertEqual([p.name for p in self.out_dir.iterdir()], [])

    def test_removed_marker_count_reports_the_dropped_markers(self) -> None:
        preview = self.preview({"raw_text": "[Verse 1]\n가사\n[Chorus]\n×2"})
        self.assertEqual(preview["lines"], ["가사"])
        self.assertEqual(preview["removed_marker_count"], 3)


class PreviewNormalizationTests(PreviewRouteTestCase):
    """The preview uses the same server normalization the build uses (M1/M2)."""

    def test_real_lyric_lines_starting_with_a_keyword_are_kept(self) -> None:
        preview = self.preview({"raw_text": "chorus line\nintro to me\noutro music"})
        self.assertEqual(
            preview["lines"], ["chorus line", "intro to me", "outro music"]
        )
        self.assertEqual(preview["removed_marker_count"], 0)
        self.assertEqual(
            preview["pairs"], [["chorus line", "intro to me"], ["outro music", None]]
        )

    def test_marker_only_forms_are_removed(self) -> None:
        preview = self.preview({"raw_text": "[Chorus]\n(후렴)\n후렴 x2\n※\n×2\n간주"})
        self.assertEqual(preview["lines"], [])
        self.assertEqual(preview["pairs"], [])
        self.assertEqual(preview["removed_marker_count"], 6)

    def test_odd_line_count_produces_a_null_tail(self) -> None:
        preview = self.preview({"raw_text": "첫 줄\n둘째 줄\n셋 줄"})
        self.assertEqual(preview["pairs"], [["첫 줄", "둘째 줄"], ["셋 줄", None]])

    def test_empty_raw_text_returns_an_empty_preview(self) -> None:
        self.assertEqual(
            self.preview({"raw_text": ""}),
            {"lines": [], "pairs": [], "removed_marker_count": 0, "slide_index": 19},
        )

    def test_dedupe_chorus_flag_reaches_the_preview(self) -> None:
        raw = "반복\nA\n반복\nB\n반복\nC"
        without = self.preview({"raw_text": raw})
        with_flag = self.preview({"raw_text": raw, "dedupe_chorus": True})

        # ``lines`` is the normalized text (before pairing); the flag only
        # changes the pairing, exactly like the build path.
        self.assertEqual(without["lines"], ["반복", "A", "반복", "B", "반복", "C"])
        self.assertEqual(with_flag["lines"], without["lines"])

        def flat(payload: dict) -> list[str]:
            return [line for pair in payload["pairs"] for line in pair if line is not None]

        self.assertEqual(flat(without), ["반복", "A", "반복", "B", "반복", "C"])
        self.assertEqual(flat(with_flag), ["반복", "A", "반복", "B", "C"])


class PreviewSlideIndexTests(PreviewRouteTestCase):
    """The chosen background slide is echoed back and validated (single-slide)."""

    def setUp(self) -> None:
        super().setUp()
        if not TEMPLATE.is_file():
            self.skipTest(f"template not found: {TEMPLATE}")

    def test_default_slide_index_is_nineteen(self) -> None:
        self.assertEqual(self.preview({"raw_text": "가\n나"})["slide_index"], 19)

    def test_explicit_slide_index_is_echoed(self) -> None:
        payload = self.preview({"raw_text": "가\n나", "slide_index": 5})
        self.assertEqual(payload["slide_index"], 5)

    def test_out_of_range_slide_index_is_a_bad_request(self) -> None:
        for bad in (0, 21, -1):
            with self.subTest(slide_index=bad):
                status, _, body = self.post_json(
                    "/api/preview", {"raw_text": "가\n나", "slide_index": bad}
                )
                self.assertEqual(status, 400)
                self.assertEqual(json.loads(body)["error"]["code"], "BAD_REQUEST")

    def test_non_integer_slide_index_is_a_bad_request(self) -> None:
        for bad in (1.5, True, "19"):
            with self.subTest(slide_index=bad):
                status, _, body = self.post_json(
                    "/api/preview", {"raw_text": "가\n나", "slide_index": bad}
                )
                self.assertEqual(status, 400)
                self.assertEqual(json.loads(body)["error"]["code"], "BAD_REQUEST")


class PreviewErrorTests(PreviewRouteTestCase):
    """Errors reuse the ``describe_error`` mapping (docs/arch/05 §3)."""

    def test_missing_raw_text_is_a_bad_request(self) -> None:
        status, _, body = self.post_json("/api/preview", {})
        self.assertEqual(status, 400)
        self.assertEqual(json.loads(body)["error"]["code"], "BAD_REQUEST")

    def test_non_string_raw_text_is_a_bad_request(self) -> None:
        for payload in ({"raw_text": 123}, {"raw_text": ["a", "b"]}, {"raw_text": None}):
            with self.subTest(payload=payload):
                status, _, body = self.post_json("/api/preview", payload)
                self.assertEqual(status, 400)
                self.assertEqual(json.loads(body)["error"]["code"], "BAD_REQUEST")

    def test_invalid_json_is_a_bad_request(self) -> None:
        status, _, body = self.request(
            "POST",
            "/api/preview",
            body=b"{not json",
            headers={"Content-Type": "application/json"},
        )
        self.assertEqual(status, 400)
        self.assertEqual(json.loads(body)["error"]["code"], "BAD_REQUEST")

    def test_non_object_json_body_is_a_bad_request(self) -> None:
        status, _, _ = self.request(
            "POST",
            "/api/preview",
            body=b"[1, 2, 3]",
            headers={"Content-Type": "application/json"},
        )
        self.assertEqual(status, 400)

    def test_oversized_body_is_rejected_with_413(self) -> None:
        connection = http.client.HTTPConnection(self.host, self.port, timeout=60)
        try:
            connection.putrequest("POST", "/api/preview")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", str(2 * 1024 * 1024))
            connection.putheader("Connection", "close")
            connection.endheaders()
            response = connection.getresponse()
            payload = response.read()
            status = response.status
        finally:
            connection.close()
        self.assertEqual(status, 413)
        self.assertEqual(json.loads(payload)["error"]["code"], "PAYLOAD_TOO_LARGE")

    def test_get_on_the_preview_route_is_404(self) -> None:
        status, _, body = self.request("GET", "/api/preview")
        self.assertEqual(status, 404)
        self.assertEqual(json.loads(body)["error"]["code"], "NOT_FOUND")

    def test_error_bodies_never_leak_a_stack_trace(self) -> None:
        _, _, body = self.post_json("/api/preview", {})
        text = body.decode("utf-8")
        self.assertNotIn("Traceback", text)
        self.assertNotIn("lyrics2pptx/", text)


if __name__ == "__main__":
    unittest.main()
